"""Подготовка деловых документов (письмо, служебная записка, справка) по записи
из базы или по данным в целом.

С моделью — связный текст, без модели — структурный шаблон с фактами и полями
[указать]. Готовый DOCX скачивается из интерфейса: на защите это сильный ход —
сервис отдаёт не картинку, а документ, который можно отправить.
"""
from __future__ import annotations

import io
from datetime import date

from sqlalchemy.orm import Session

from ..models import Record, User
from . import assistant, client
from .prompts import ARTIFACT_KINDS, SYSTEM_ARTIFACT


def _context_for(db: Session, record_id: int | None) -> tuple[str, str]:
    """Возвращает (факты, как называть предмет документа)."""
    if record_id:
        rec = db.get(Record, record_id)
        if rec is None:
            raise ValueError(f"Запись id={record_id} не найдена")
        return assistant.record_facts(rec), f"записи {rec.key}"
    st = assistant.project_stats(db)
    top = db.query(Record).order_by(Record.amount.desc().nullslast()).limit(7).all()
    ctx = (
        f"Всего записей: {st['total']}; сумма по полю amount: {assistant.fmt_num(st['amount_sum'])}.\n\n"
        + "\n".join(assistant.record_facts(r) for r in top)
    )
    return ctx, "данным в целом"


def _fallback_text(kind: str, context: str, subject_label: str) -> str:
    today = date.today().strftime("%d.%m.%Y")
    head = ARTIFACT_KINDS[kind]["title"]
    common = (
        f"{head} по {subject_label}\nДата: {today}\n\n"
        "(Языковая модель недоступна — сформирован структурный шаблон с фактами; "
        "заполните поля [указать] вручную.)\n\n"
    )
    if kind == "letter":
        body = ("Кому: [указать: адресат]\nТема: [указать: тема]\n\n"
                "Уважаемые коллеги!\n\nПо имеющимся данным фиксируем следующее:\n\n"
                f"{context}\n\nПросим в срок до [указать: дата] представить пояснения.\n\n"
                "С уважением,\n[указать: должность, ФИО]")
    elif kind == "memo":
        body = ("Кому: [указать: адресат]\nОт кого: [указать: подразделение]\n"
                "Тема: [указать: тема]\n\nИзложение ситуации:\n\n" + context +
                "\n\nПредлагаемые решения:\n1. [указать]\n2. [указать]\n\n"
                "Прошу принять решение по пункту [указать].")
    else:  # summary
        body = "СПРАВКА\n\nСостояние дел:\n\n" + context + "\n\nВыводы и рекомендации: [указать]."
    return common + body


def generate(db: Session, user: User, kind: str, record_id: int | None = None,
             instructions: str = "") -> dict:
    if kind not in ARTIFACT_KINDS:
        raise ValueError(f"Неизвестный вид документа: {kind}")
    context, label = _context_for(db, record_id)
    spec = ARTIFACT_KINDS[kind]
    llm_used = False
    if client.llm_available():
        try:
            extra = f"\n\nДополнительные указания пользователя: {instructions}" if instructions else ""
            text = client.chat([
                {"role": "system", "content": SYSTEM_ARTIFACT},
                {"role": "user", "content": (
                    f"{spec['instruction']}{extra}\n\nРечь идёт о {label}.\n\nФАКТЫ:\n{context}"
                )},
            ], temperature=0.3, max_tokens=1800)
            llm_used = True
        except Exception:
            text = _fallback_text(kind, context, label)
    else:
        text = _fallback_text(kind, context, label)
    return {"kind": kind, "title": spec["title"], "text": text, "llm_used": llm_used}


def to_docx(title: str, text: str) -> io.BytesIO:
    from docx import Document as DocxDocument

    doc = DocxDocument()
    doc.add_heading(title, level=1)
    for block in text.split("\n\n"):
        for line in block.split("\n"):
            stripped = line.strip()
            if not stripped:
                continue
            if stripped.startswith(("- ", "• ", "* ")):
                doc.add_paragraph(stripped[2:].strip(), style="List Bullet")
            else:
                doc.add_paragraph(stripped)
        doc.add_paragraph("")
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf

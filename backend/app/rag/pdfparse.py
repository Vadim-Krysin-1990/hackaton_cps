"""Парсинг документов базы знаний: PDF (PyMuPDF), DOCX, TXT/MD.
Возвращает постраничный текст — номер страницы едет в Evidence ответа.
"""
from __future__ import annotations

import re
from pathlib import Path


def parse_pdf(path: Path) -> list[tuple[int, str]]:
    import fitz  # PyMuPDF

    pages: list[tuple[int, str]] = []
    with fitz.open(path) as doc:
        for i, page in enumerate(doc, start=1):
            text = page.get_text("text").strip()
            if text:
                pages.append((i, text))
    return pages


def _docx_text_from_xml(path: Path) -> str:
    """Запасной разбор: часть .docx с сайта собрана так, что python-docx падает
    на незарегистрированных вложениях (SVG). Текст берём прямо из document.xml.
    """
    import zipfile
    from xml.etree import ElementTree

    ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    with zipfile.ZipFile(path) as z:
        xml = z.read("word/document.xml")
    root = ElementTree.fromstring(xml)
    lines = []
    for para in root.iter(f"{ns}p"):
        text = "".join(node.text or "" for node in para.iter(f"{ns}t"))
        if text.strip():
            lines.append(text.strip())
    return "\n".join(lines)


def parse_docx(path: Path) -> list[tuple[int, str]]:
    try:
        from docx import Document as DocxDocument

        doc = DocxDocument(str(path))
        text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    except Exception:
        text = _docx_text_from_xml(path)
    return [(1, text)] if text.strip() else []


def parse_text(path: Path) -> list[tuple[int, str]]:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "utf-8", "cp1251"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        text = raw.decode("utf-8", errors="replace")
    # сохранённые страницы сайта чистятся тем же фильтром, что и HTML
    lines = [ln for ln in text.splitlines()
             if ln.strip().lower() not in _BOILERPLATE and not _READ_TIME.match(ln.strip())]
    text = "\n".join(lines)
    return [(1, text)] if text.strip() else []


_MAIN_CLASS = re.compile(r"content|main|inner|page|article|text|book", re.I)

# служебные надписи шаблона сайта — в поиске только мешают
_BOILERPLATE = {
    "данный функционал доступен после авторизации", "печать", "поделиться",
    "на чтение", "руководство пользователя", "глоссарий", "разделы",
    "главная страница", "главная", "версия для печати", "наверх",
}
_READ_TIME = re.compile(r"^\d+\s+минут[аы]?$|^\d+$")


def html_to_text(html: bytes | str) -> str:
    """Текст страницы без обвязки. Основной блок выбирается по объёму текста:
    на Bitrix-вёрстке первый попавшийся div.content — это хлебные крошки.
    """
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(html, "html.parser")
    for tag in soup(["script", "style", "noscript", "svg", "iframe", "form",
                     "nav", "header", "footer", "aside"]):
        tag.decompose()

    candidates = soup.find_all(["main", "article"]) + soup.find_all(class_=_MAIN_CLASS)
    best, best_len = None, 0
    for node in candidates:
        length = len(node.get_text(" ", strip=True))
        if length > best_len:
            best, best_len = node, length
    root = best if best_len >= 400 else (soup.body or soup)

    lines, prev = [], None
    for chunk in root.get_text("\n").splitlines():
        chunk = re.sub(r"[ \t\xa0]+", " ", chunk).strip()
        if not chunk or chunk == prev:
            prev = chunk
            continue
        if chunk.lower() in _BOILERPLATE or _READ_TIME.match(chunk):
            continue
        lines.append(chunk)
        prev = chunk
    return "\n".join(lines)


def parse_html(path: Path) -> list[tuple[int, str]]:
    text = html_to_text(path.read_bytes())
    return [(1, text)] if text.strip() else []


def parse_xlsx(path: Path) -> list[tuple[int, str]]:
    """Лист книги -> «страница»: строки склеиваются в текст, чтобы перечни
    (группы МТР, процедуры) попадали в поиск."""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    pages: list[tuple[int, str]] = []
    for i, ws in enumerate(wb.worksheets, start=1):
        lines = [ws.title]
        for row in ws.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                lines.append(" | ".join(cells))
        text = "\n".join(lines)
        if len(text) > 50:
            pages.append((i, text))
    wb.close()
    return pages


def parse_document(path: str | Path) -> list[tuple[int, str]]:
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".pdf":
        return parse_pdf(path)
    if ext in (".docx", ".docm"):
        return parse_docx(path)
    if ext in (".txt", ".md"):
        return parse_text(path)
    if ext in (".html", ".htm"):
        return parse_html(path)
    if ext in (".xlsx", ".xlsm"):
        return parse_xlsx(path)
    raise ValueError(
        f"Неподдерживаемый формат документа: {ext} "
        f"(поддержаны pdf, docx, xlsx, txt, md, html)")


# служебные надписи титульного листа — заголовком документа не являются
_TITLE_NOISE = re.compile(
    r"^(утвержд|приложение|согласовано|прилож\.|приказ|стр\.|лист|copy|конфиденциально|"
    r"[\W\d]+)$|^\d",
    re.IGNORECASE,
)


def guess_title(pages: list[tuple[int, str]], fallback: str) -> str:
    """Заголовок = самая содержательная строка первой страницы (не «УТВЕРЖДЕНО» и не номер)."""
    if not pages:
        return fallback
    candidates = []
    for line in pages[0][1].splitlines()[:40]:
        line = re.sub(r"\s+", " ", line).strip(" .:-—")
        if len(line) < 12 or _TITLE_NOISE.match(line):
            continue
        letters = sum(ch.isalpha() for ch in line)
        if letters < len(line) * 0.6:
            continue
        candidates.append(line)
    if not candidates:
        return fallback
    return max(candidates[:12], key=len)[:200]

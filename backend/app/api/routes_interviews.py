"""Exit-интервью: загрузка транскриптов, паспорта, дашборд, выгрузка."""
from __future__ import annotations

import io
import json
import time
from collections import Counter
from itertools import combinations
from pathlib import Path

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from sqlalchemy import func
from sqlalchemy.orm import Session

from ..config import settings
from ..db import get_db
from ..domain.exit_interview import pipeline
from ..llm import client as llm_client
from ..models import Insight, Interview, User
from .deps import audit, get_current_user, timed

router = APIRouter(prefix="/interviews", tags=["interviews"])

ALLOWED_EXT = {".txt", ".md"}
MAX_CHARS = 20_000
MIN_CHARS = 80


def _decode(raw: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "cp1251"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _row(i: Interview) -> dict:
    p = i.passport or {}
    return {
        "id": i.id, "filename": i.filename, "chars": i.chars, "status": i.status, "error": i.error,
        "exit_reason": i.exit_reason, "risk_zone": i.risk_zone, "top_problem": i.top_problem,
        "department": i.department or "", "manager": i.manager or "",
        "engine": i.engine, "duration_ms": i.duration_ms,
        "summary": p.get("summary", ""), "pains": len(p.get("pain_points", [])),
        "created_at": i.created_at.isoformat() if i.created_at else None,
    }


def _analyze(db: Session, user: User, filename: str, text: str, model: str | None) -> Interview:
    text = text.strip()
    if len(text) < MIN_CHARS:
        raise HTTPException(status_code=400, detail=f"Текст слишком короткий ({len(text)} символов): "
                            "это не похоже на транскрипт беседы.")
    if len(text) > MAX_CHARS:
        raise HTTPException(status_code=400, detail=f"Текст длиннее {MAX_CHARS} символов. "
                            "Разбейте беседу на части.")
    rec = Interview(filename=filename, source=text, chars=len(text), uploaded_by=user.id)
    db.add(rec)
    db.commit()
    try:
        with timed() as t:
            result = pipeline.run(text, model=model)
    except Exception as e:  # сюда попадает только ошибка самого конвейера, не модели
        rec.status, rec.error = "error", f"{type(e).__name__}: {e}"
        db.commit()
        audit(db, user, "interview_failed", {"filename": filename, "error": rec.error})
        raise HTTPException(status_code=422, detail=f"Не удалось разобрать транскрипт: {e}")
    _apply(rec, result, t.ms)
    db.commit()
    audit(db, user, "interview_analyze", {
        "id": rec.id, "filename": filename, "engine": rec.engine,
        "quotes_checked": result["verification"]["checked"],
        "quotes_rejected": len(result["verification"]["rejected"]),
        "steps": [f"{s['step']} {s['ms']} мс" for s in result["steps"]],
    }, duration_ms=t.ms)
    return rec


def _apply(rec: Interview, result: dict, ms: int) -> None:
    p = result["passport"]
    rec.passport = p
    rec.analysis = {k: result[k] for k in ("sentiment", "verification", "utterances", "steps", "notes")}
    rec.exit_reason, rec.risk_zone = p.get("exit_reason"), p.get("risk_zone")
    rec.top_problem, rec.engine = p.get("top_problem"), result["engine"]
    rec.department, rec.manager = (p.get("department") or None), (p.get("manager") or None)
    rec.duration_ms, rec.status, rec.error = ms, "done", None


@router.get("/models")
def models(user: User = Depends(get_current_user)):
    """Переключатель модели: что доступно у провайдера и что выбрано по умолчанию."""
    return {"default": settings.llm_model, "available": llm_client.llm_available(),
            "provider": llm_client.provider(), "models": llm_client.list_models()}


@router.post("")
def upload(files: list[UploadFile] = File(...), model: str = Form(""),
           db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    """Один или несколько .txt за раз: датасет организаторов грузится одним действием."""
    out, errors = [], []
    for f in files:
        ext = Path(f.filename or "").suffix.lower()
        if ext not in ALLOWED_EXT:
            errors.append({"filename": f.filename, "error": f"Формат {ext or 'без расширения'} не поддерживается, нужен .txt"})
            continue
        try:
            rec = _analyze(db, user, f.filename or "transcript.txt", _decode(f.file.read()), model or None)
            out.append(_row(rec))
        except HTTPException as e:
            errors.append({"filename": f.filename, "error": e.detail})
    if not out and errors:
        raise HTTPException(status_code=422, detail=errors[0]["error"])
    return {"items": out, "errors": errors}


class TextIn(BaseModel):
    text: str
    filename: str = "вставленный текст"
    model: str = ""


@router.post("/text")
def analyze_text(body: TextIn, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    rec = _analyze(db, user, body.filename, body.text, body.model or None)
    return _full(rec)


@router.get("")
def list_interviews(q: str = "", reason: str = "", risk: str = "", department: str = "", manager: str = "",
                    limit: int = Query(50, ge=1, le=500), offset: int = Query(0, ge=0),
                    db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    query = db.query(Interview)
    if q.strip():
        like = f"%{q.strip().lower()}%"
        query = query.filter(func.lower(Interview.source).like(like) | func.lower(Interview.filename).like(like))
    if reason:
        query = query.filter(Interview.exit_reason == reason)
    if risk:
        query = query.filter(Interview.risk_zone == risk)
    if department:
        query = query.filter(Interview.department == department)
    if manager:
        query = query.filter(Interview.manager == manager)
    total = query.count()
    rows = query.order_by(Interview.id.desc()).offset(offset).limit(limit).all()
    return {"items": [_row(r) for r in rows], "total": total}


def _full(rec: Interview) -> dict:
    return {**_row(rec), "source": rec.source, "passport": rec.passport or {},
            "analysis": rec.analysis or {}}


@router.get("/dashboard")
def dashboard(db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    """Сводка по всем интервью: то, что видит руководитель.

    Считается из паспортов, не из модели: цифры одинаковы при любом движке."""
    rows = db.query(Interview).filter(Interview.status == "done").all()
    reasons, risks, engines = Counter(), Counter(), Counter()
    pain_interviews, pain_mentions = Counter(), Counter()
    practices = Counter()
    edges = Counter()
    checked = accepted = rejected = 0
    durations = []
    latest_suggestions: dict[str, list[str]] = {}
    for r in rows:
        p = r.passport or {}
        reasons[p.get("exit_reason") or "другое"] += 1
        risks[p.get("risk_zone") or "Medium"] += 1
        engines[r.engine or "—"] += 1
        durations.append(r.duration_ms or 0)
        cats = sorted({pp.get("category") or "другое" for pp in p.get("pain_points", [])})
        for pp in p.get("pain_points", []):
            c = pp.get("category") or "другое"
            pain_mentions[c] += int(pp.get("mentions") or 1)
        for c in cats:
            pain_interviews[c] += 1
        for a, b in combinations(cats, 2):
            edges[(a, b)] += 1
        for bp in p.get("best_practices", []):
            practices[bp.get("label") or "—"] += 1
        v = (r.analysis or {}).get("verification") or {}
        checked += v.get("checked", 0)
        accepted += v.get("accepted", 0)
        rejected += len(v.get("rejected", []))
        tp = p.get("top_problem")
        if tp and tp not in latest_suggestions:
            latest_suggestions[tp] = p.get("improvement_suggestions", [])[:3]
    top_problem = pain_interviews.most_common(1)[0][0] if pain_interviews else None
    return {
        "total": len(rows),
        "by_reason": [{"name": k, "count": v} for k, v in reasons.most_common()],
        "by_risk": [{"name": k, "count": v} for k, v in risks.most_common()],
        "by_engine": [{"name": k, "count": v} for k, v in engines.most_common()],
        "pains": [{"name": k, "count": v, "mentions": pain_mentions[k],
                   "share": round(v / len(rows) * 100) if rows else 0}
                  for k, v in pain_interviews.most_common()],
        "practices": [{"name": k, "count": v} for k, v in practices.most_common(8)],
        "edges": [{"a": a, "b": b, "count": c} for (a, b), c in edges.most_common(30)],
        "top_problem": top_problem,
        "top_suggestions": latest_suggestions.get(top_problem, []) if top_problem else [],
        "verification": {"checked": checked, "accepted": accepted, "rejected": rejected},
        "avg_duration_ms": int(sum(durations) / len(durations)) if durations else 0,
        "manual_minutes": settings.baseline_manual_minutes,
    }



@router.get("/graph")
def graph(db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    """Клубок связей: интервью ↔ отдел, руководитель, причина, проблема.

    Узлы пяти типов, рёбра только через интервью: так видно, из какого отдела
    и по какой причине уходят, и какие проблемы тянутся за одним руководителем."""
    rows = db.query(Interview).filter(Interview.status == "done").all()
    nodes: dict[str, dict] = {}
    edges: Counter = Counter()

    def node(kind: str, name: str) -> str:
        key = f"{kind}:{name}"
        if key not in nodes:
            nodes[key] = {"id": key, "kind": kind, "name": name, "count": 0}
        nodes[key]["count"] += 1
        return key

    for r in rows:
        p = r.passport or {}
        iv = node("interview", r.filename)
        nodes[iv].update({"risk": r.risk_zone, "summary": p.get("summary", ""), "interview_id": r.id})
        links = [node("reason", p.get("exit_reason") or "другое")]
        if r.department:
            links.append(node("department", r.department))
        if r.manager:
            links.append(node("manager", r.manager))
        for pp in p.get("pain_points", []):
            links.append(node("problem", pp.get("category") or "другое"))
        for k in dict.fromkeys(links):
            edges[(iv, k)] += 1
    return {
        "nodes": list(nodes.values()),
        "edges": [{"source": a, "target": b, "count": c} for (a, b), c in edges.items()],
        "kinds": {"interview": "интервью", "department": "отдел", "manager": "руководитель",
                  "reason": "причина ухода", "problem": "проблема"},
    }


def _insight_context(db: Session, problem: str = "") -> tuple[str, int]:
    """Сводка по всем паспортам — то, что видит модель. Только агрегаты и цитаты.
    problem — категория, на которой нужно сосредоточиться: её цитаты идут первыми."""
    rows = db.query(Interview).filter(Interview.status == "done").all()
    if not rows:
        return "", 0
    focus_quotes, focus_where = [], Counter()
    reasons, pains, depts, mgrs, risks = Counter(), Counter(), Counter(), Counter(), Counter()
    dept_reason, mgr_reason = Counter(), Counter()
    quotes = []
    for r in rows:
        p = r.passport or {}
        reasons[p.get("exit_reason") or "другое"] += 1
        risks[p.get("risk_zone") or "Medium"] += 1
        if r.department:
            depts[r.department] += 1
            dept_reason[(r.department, p.get("exit_reason") or "другое")] += 1
        if r.manager:
            mgrs[r.manager] += 1
            mgr_reason[(r.manager, p.get("exit_reason") or "другое")] += 1
        for pp in p.get("pain_points", []):
            cat = pp.get("category") or "другое"
            pains[cat] += 1
            if problem and cat == problem:
                focus_where[(r.department or "отдел не указан", r.manager or "руководитель не указан")] += 1
                for q in pp.get("quotes", [])[:2]:
                    if len(focus_quotes) < 10:
                        focus_quotes.append(f"«{q}»")
            elif pp.get("quotes") and len(quotes) < 12:
                quotes.append(f"[{cat}] «{pp['quotes'][0]}»")
    lines = [f"Интервью всего: {len(rows)}",
             "Причины ухода: " + ", ".join(f"{k} — {v}" for k, v in reasons.most_common()),
             "Системные проблемы (в скольких интервью): " + ", ".join(f"{k} — {v}" for k, v in pains.most_common()),
             "Риск для команды: " + ", ".join(f"{k} — {v}" for k, v in risks.most_common())]
    if depts:
        lines.append("Отделы: " + ", ".join(f"{k} — {v}" for k, v in depts.most_common()))
        lines.append("Отдел и причина: " + ", ".join(f"{d} / {r} — {v}" for (d, r), v in dept_reason.most_common(8)))
    if mgrs:
        lines.append("Руководители: " + ", ".join(f"{k} — {v}" for k, v in mgrs.most_common()))
        lines.append("Руководитель и причина: " + ", ".join(f"{m} / {r} — {v}" for (m, r), v in mgr_reason.most_common(8)))
    if problem:
        lines.append(f"ФОКУС: проблема «{problem}» названа в {pains.get(problem, 0)} интервью из {len(rows)}. "
                     "Где встречается (отдел / руководитель — сколько раз): "
                     + (", ".join(f"{d} / {m} — {c}" for (d, m), c in focus_where.most_common(6)) or "нет данных"))
        lines.append("Цитаты по этой проблеме:\n" + ("\n".join(focus_quotes) or "цитат нет"))
    lines.append("Характерные цитаты по остальным проблемам:\n" + "\n".join(quotes))
    return "\n".join(lines), len(rows)


def _insight_payload(i: Insight | None) -> dict:
    if i is None:
        return {"ready": False}
    return {"ready": True, "problem": i.problem, "interviews_count": i.interviews_count, "engine": i.engine,
            "duration_ms": i.duration_ms, "created_at": i.created_at.isoformat() if i.created_at else None,
            **(i.payload or {})}


def _build_insight(db: Session, user: User, model: str | None, problem: str = "") -> Insight:
    from ..domain.exit_interview.pipeline import _extract_json, load_prompt

    context, n = _insight_context(db, problem)
    if not n:
        raise HTTPException(status_code=400, detail="Нет разобранных интервью — вывод строить не из чего")
    system = load_prompt("insight_system")
    engine = "fallback"
    payload: dict | None = None
    t0 = time.perf_counter()
    if llm_client.llm_available():
        try:
            raw = llm_client.chat([{"role": "system", "content": system},
                                   {"role": "user", "content": "Сводка по интервью:\n\n" + context +
                                    (f"\n\nРекомендации нужны ТОЛЬКО по проблеме «{problem}»: что с ней делать, "
                                     "где она сконцентрирована, кто отвечает." if problem else "") +
                                    "\n\nВерни вывод для руководителя в формате JSON."}],
                                  temperature=0.2, max_tokens=2500, model=model, json_mode=True)
            data = _extract_json(raw)
            payload = {
                "headline": str(data.get("headline") or ""),
                "summary": str(data.get("summary") or ""),
                "signals": [str(x) for x in (data.get("signals") or [])][:5],
                "recommendations": [
                    {"action": str(r.get("action") or ""), "owner": str(r.get("owner") or ""),
                     "effect": str(r.get("effect") or "")}
                    for r in (data.get("recommendations") or []) if isinstance(r, dict)][:5],
            }
            engine = f"llm:{model or settings.llm_model}"
        except Exception as e:
            payload = None
            engine = f"fallback ({type(e).__name__})"
    if payload is None:
        from ..domain.exit_interview import fallback as fb
        dash = dashboard(db, user)
        top = problem or dash["top_problem"] or ""
        payload = {"headline": f"Проблема: {top or 'не выделена'}",
                   "summary": f"Разобрано {n} интервью. Модель недоступна, вывод собран из агрегатов.",
                   "signals": [], "recommendations": [{"action": s, "owner": "HR", "effect": ""}
                                                       for s in fb.SUGGESTIONS.get(top, fb.DEFAULT_SUGGESTIONS)]}
    ins = Insight(problem=problem, interviews_count=n, engine=engine, payload=payload,
                  duration_ms=int((time.perf_counter() - t0) * 1000))
    db.add(ins)
    db.commit()
    audit(db, user, "insight_build", {"interviews": n, "engine": engine, "problem": problem}, duration_ms=ins.duration_ms)
    return ins


@router.get("/insight")
def insight(refresh: bool = False, model: str = "", problem: str = "", db: Session = Depends(get_db),
            user: User = Depends(get_current_user)):
    """Вывод ИИ для дашборда: по главной проблеме или по выбранной (problem).
    Кэш на каждую проблему свой; пересобирается, если добавились интервью или по кнопке."""
    n = db.query(func.count(Interview.id)).filter(Interview.status == "done").scalar() or 0
    last = db.query(Insight).filter(Insight.problem == problem).order_by(Insight.id.desc()).first()
    if n == 0:
        return {"ready": False}
    if refresh or last is None or last.interviews_count != n:
        last = _build_insight(db, user, model or None, problem)
    return _insight_payload(last)


@router.get("/{interview_id}")
def get_interview(interview_id: int, db: Session = Depends(get_db),
                  user: User = Depends(get_current_user)):
    rec = db.get(Interview, interview_id)
    if rec is None:
        raise HTTPException(status_code=404, detail="Интервью не найдено")
    audit(db, user, "interview_view", {"id": interview_id})
    return _full(rec)


@router.post("/{interview_id}/reanalyze")
def reanalyze(interview_id: int, model: str = Query(""), db: Session = Depends(get_db),
              user: User = Depends(get_current_user)):
    """Тот же текст другой моделью: сравнение малой и большой модели на одном интервью."""
    rec = db.get(Interview, interview_id)
    if rec is None:
        raise HTTPException(status_code=404, detail="Интервью не найдено")
    with timed() as t:
        result = pipeline.run(rec.source, model=model or None)
    _apply(rec, result, t.ms)
    db.commit()
    audit(db, user, "interview_analyze", {"id": rec.id, "filename": rec.filename, "engine": rec.engine,
                                          "reanalyze": True,
                                          "quotes_rejected": len(result["verification"]["rejected"])},
          duration_ms=t.ms)
    return _full(rec)


@router.delete("/{interview_id}")
def delete_interview(interview_id: int, db: Session = Depends(get_db),
                     user: User = Depends(get_current_user)):
    rec = db.get(Interview, interview_id)
    if rec is None:
        raise HTTPException(status_code=404, detail="Интервью не найдено")
    db.delete(rec)
    db.commit()
    audit(db, user, "interview_delete", {"id": interview_id})
    return {"ok": True}


@router.get("/{interview_id}/docx")
def export_docx(interview_id: int, db: Session = Depends(get_db),
                user: User = Depends(get_current_user)):
    from docx import Document as DocxDocument

    rec = db.get(Interview, interview_id)
    if rec is None or not rec.passport:
        raise HTTPException(status_code=404, detail="Паспорт не найден")
    p = rec.passport
    doc = DocxDocument()
    doc.add_heading(f"Паспорт проблемы: {rec.filename}", level=1)
    doc.add_paragraph(p.get("summary", ""))
    doc.add_heading("Причина ухода", level=2)
    doc.add_paragraph(f"{p.get('exit_reason', '')}. «{p.get('exit_reason_quote', '')}»")
    doc.add_paragraph(f"Говорит: {p.get('says', '')}")
    doc.add_paragraph(f"Чувствует: {p.get('feels', '')}")
    doc.add_heading("Системные проблемы", level=2)
    for pp in p.get("pain_points", []):
        doc.add_paragraph(f"{pp['label']} ({pp['category']}, упоминаний: {pp['mentions']})", style="List Bullet")
        for q in pp.get("quotes", []):
            doc.add_paragraph(f"«{q}»", style="List Bullet 2")
    doc.add_heading("Что работает хорошо", level=2)
    for bp in p.get("best_practices", []):
        doc.add_paragraph(f"{bp['label']}: «{bp['quote']}»", style="List Bullet")
    doc.add_heading("Зона риска", level=2)
    doc.add_paragraph(f"{p.get('risk_zone', '')}. {p.get('risk_rationale', '')}")
    doc.add_heading("Гипотезы по главной проблеме", level=2)
    for s in p.get("improvement_suggestions", []):
        doc.add_paragraph(s, style="List Number")
    v = (rec.analysis or {}).get("verification", {})
    doc.add_paragraph(f"Проверка цитат: проверено {v.get('checked', 0)}, подтверждено {v.get('accepted', 0)}, "
                      f"отклонено {len(v.get('rejected', []))}. Движок: {rec.engine}.")
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    audit(db, user, "interview_export", {"id": rec.id})
    name = Path(rec.filename).stem
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''passport_{interview_id}.docx"},
    )
